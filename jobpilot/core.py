from __future__ import annotations
import hashlib
import json
import os
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

STATUSES = {'DISCOVERED','PREPARED','FILLING','NEEDS_REVIEW','BLOCKED','SUBMITTING',
            'SUBMISSION_UNCONFIRMED','SUBMITTED','CONFIRMED','INTERVIEW','REJECTED','OFFER','SKIPPED'}
NO_RETRY = {'SUBMITTING','SUBMISSION_UNCONFIRMED','SUBMITTED','CONFIRMED','INTERVIEW','REJECTED','OFFER'}

def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')

def canonical_url(url: str) -> str:
    p=urlsplit(url.strip())
    if p.scheme not in ('http','https') or not p.hostname or p.username or p.password:
        raise ValueError('Job URL must be an http(s) URL without embedded credentials.')
    query=[(k,v) for k,v in parse_qsl(p.query,keep_blank_values=True)
           if not k.lower().startswith('utm_') and k.lower() not in {'fbclid','gclid'}]
    return urlunsplit((p.scheme.lower(),p.netloc.lower(),p.path or '/',urlencode(sorted(query)),''))

def origin(url):
    p=urlsplit(url); return f'{p.scheme.lower()}://{p.netloc.lower()}'

def read_json(path): return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def write_json(path, value):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_suffix(path.suffix+'.tmp')
    tmp.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8'); os.replace(tmp,path)

def validate_profile(p):
    if not isinstance(p,dict) or p.get('facts_verified') is not True:
        raise ValueError('Review your profile, check that the details are accurate, and save it as verified before preparing applications.')
    personal=p.get('personal',{})
    def text(value): return isinstance(value,str) and bool(value.strip())
    def texts(value): return isinstance(value,list) and all(text(v) for v in value)
    if not isinstance(personal,dict): raise ValueError('personal must be an object.')
    missing=[k for k in ('first_name','last_name','email','phone','city','country') if not text(personal.get(k))]
    if missing: raise ValueError('Missing profile fields: '+', '.join(missing))
    if '@' not in personal['email']: raise ValueError('Invalid email address.')
    if not texts(p.get('skills')) or not p['skills']:
        raise ValueError('skills must be a nonempty array of strings.')
    for section in ('summary_variants','languages'):
        if not texts(p.get(section,[])): raise ValueError(section+' must be an array of strings.')
    if not isinstance(p.get('summary',''),str): raise ValueError('summary must be a string.')
    if not isinstance(p.get('answers',{}),dict): raise ValueError('answers must be an object.')
    for section,keys in [('experience',('company','role','start','end')),('projects',('name',)),('education',('degree','institution','dates'))]:
        entries=p.get(section,[])
        if not isinstance(entries,list): raise ValueError(section+' must be an array.')
        for entry in entries:
            if not isinstance(entry,dict) or any(not text(entry.get(k)) for k in keys):
                raise ValueError(section+' entries require text fields: '+', '.join(keys))
            if not texts(entry.get('bullets',[])) or (section!='education' and not entry.get('bullets')):
                raise ValueError(section+' bullets must be an array of factual statements.')
    if not (p.get('experience') or p.get('projects')):
        raise ValueError('Add your real skills and at least one work experience or project.')
    return p

def excel_safe(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')): return "'"+value
    return value

class Store:
    def __init__(self, home):
        self.home=Path(home); self.home.mkdir(parents=True,exist_ok=True)
        self.path=self.home/'applications.sqlite3'
        with self.connect() as c:
            c.executescript('''CREATE TABLE IF NOT EXISTS jobs (
                id TEXT PRIMARY KEY, url TEXT UNIQUE NOT NULL, title TEXT, company TEXT,
                location TEXT, description TEXT, source TEXT, created TEXT, updated TEXT,
                status TEXT, resume TEXT DEFAULT '', note TEXT DEFAULT '', evidence TEXT DEFAULT '');
                CREATE TABLE IF NOT EXISTS events (id INTEGER PRIMARY KEY AUTOINCREMENT,
                job_id TEXT, created TEXT, status TEXT, note TEXT);
                CREATE TABLE IF NOT EXISTS mail_links (message_id TEXT PRIMARY KEY,job_id TEXT,created TEXT);''')
    @contextmanager
    def connect(self):
        c=sqlite3.connect(self.path,timeout=30); c.row_factory=sqlite3.Row
        try:
            with c:
                yield c
        finally:
            c.close()
    def add(self, job):
        url=canonical_url(job['url']); key=hashlib.sha256(url.encode()).hexdigest()[:16]; stamp=now()
        with self.connect() as c:
            # Re-import only refreshes job data, never resets application progress.
            old=c.execute('SELECT * FROM jobs WHERE id=?',(key,)).fetchone()
            if old:
                fields=[k for k in ('title','company','location','description','source')
                        if isinstance(job.get(k),str) and job[k].strip() and job[k]!=old[k]]
                if fields:
                    c.execute('UPDATE jobs SET '+','.join(k+'=?' for k in fields)+',updated=? WHERE id=?',
                              [job[k] for k in fields]+[stamp,key])
                    c.execute('INSERT INTO events(job_id,created,status,note) VALUES(?,?,?,?)',
                              (key,stamp,old['status'],'Imported metadata updated; re-prepare resume if requirements changed'))
                return key
            c.execute('INSERT INTO jobs (id,url,title,company,location,description,source,created,updated,status) VALUES (?,?,?,?,?,?,?,?,?,?)',
                (key,url,job.get('title','Untitled'),job.get('company','Unknown'),job.get('location',''),job.get('description',''),job.get('source','import'),stamp,stamp,'DISCOVERED'))
            c.execute('INSERT INTO events(job_id,created,status,note) VALUES(?,?,?,?)',(key,stamp,'DISCOVERED','Imported'))
        return key
    def get(self,key):
        with self.connect() as c: row=c.execute('SELECT * FROM jobs WHERE id=?',(key,)).fetchone()
        if not row: raise ValueError(f'Unknown job id: {key}')
        return dict(row)
    def rows(self):
        with self.connect() as c: return [dict(r) for r in c.execute('SELECT * FROM jobs ORDER BY created DESC,id')]
    def can_apply(self,key): return self.get(key)['status'] not in NO_RETRY|{'SKIPPED','FILLING'}
    def claim(self,key):
        with self.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            row=c.execute('SELECT status FROM jobs WHERE id=?',(key,)).fetchone()
            if not row or row['status'] in NO_RETRY|{'SKIPPED','FILLING'}:
                raise ValueError('Job is already active/submitted/uncertain/skipped. Check its history first.')
            c.execute('UPDATE jobs SET status=?,updated=? WHERE id=?',('FILLING',now(),key))
            c.execute('INSERT INTO events(job_id,created,status,note) VALUES(?,?,?,?)',(key,now(),'FILLING','Browser session claimed'))
    def update(self,key,status,note='',resume=None,evidence=None):
        if status not in STATUSES: raise ValueError('Unknown status')
        self.get(key)
        sets=['status=?','updated=?','note=?']; values=[status,now(),note]
        for field,value in [('resume',resume),('evidence',evidence)]:
            if value is not None: sets.append(field+'=?'); values.append(str(value))
        with self.connect() as c:
            c.execute('UPDATE jobs SET '+','.join(sets)+' WHERE id=?',values+[key])
            c.execute('INSERT INTO events(job_id,created,status,note) VALUES(?,?,?,?)',(key,now(),status,note))
    def export(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.table import Table,TableStyleInfo
        wb=Workbook(); ws=wb.active; ws.title='Applications'
        columns=['id','company','title','location','status','created','updated','url','resume','note','evidence']
        ws.append([x.replace('_',' ').title() for x in columns])
        for row in self.rows():
            vals=[]
            for k in columns:
                value=row[k]
                if k in ('created','updated'): value=datetime.fromisoformat(value).replace(tzinfo=None)
                vals.append(excel_safe(value))
            ws.append(vals)
        with self.connect() as c: events=c.execute('SELECT job_id,created,status,note FROM events ORDER BY id').fetchall()
        hist=wb.create_sheet('History'); hist.append(['Job ID','Timestamp (UTC)','Status','Note'])
        for row in events: hist.append([excel_safe(row[0]),datetime.fromisoformat(row[1]).replace(tzinfo=None),row[2],excel_safe(row[3])])
        for sheet,widths in [(ws,[22,25,42,25,28,22,22,55,45,55,55]),(hist,[22,24,28,80])]:
            sheet.freeze_panes='A2'; sheet.sheet_view.showGridLines=False
            for cell in sheet[1]: cell.fill=PatternFill('solid',fgColor='17365D'); cell.font=Font(color='FFFFFF',bold=True)
            sheet.row_dimensions[1].height=27
            for i,width in enumerate(widths,1): sheet.column_dimensions[sheet.cell(1,i).column_letter].width=width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
                    if isinstance(cell.value,datetime): cell.number_format='yyyy-mm-dd hh:mm'
            if sheet.max_row>1:
                table=Table(displayName=sheet.title+'Table',ref=sheet.dimensions)
                table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
                sheet.add_table(table)
        path=self.home/'applications.xlsx'; tmp=self.home/('applications.'+uuid.uuid4().hex+'.tmp.xlsx')
        wb.save(tmp)
        try: os.replace(tmp,path)
        except PermissionError:
            raise PermissionError('Close applications.xlsx in Excel and run export again. SQLite data is safe.')
        return path
