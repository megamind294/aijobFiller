import json
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
from jobpilot.core import Store,read_json
from jobpilot.sources import read_jobs,discover
from jobpilot.demo import make_server
from jobpilot.resume import prepare

PROFILE={'facts_verified':True,'personal':{'first_name':'Demo','last_name':'Candidate','email':'demo@example.com','phone':'+48123456789','city':'Krakow','country':'Poland'},'summary':'Frontend developer.','skills':['React','TypeScript'], 'experience':[{'company':'Demo Employer','role':'Developer','start':'2022','end':'2024','bullets':['Built React forms.']}],'education':[], 'projects':[],'answers':{}}

class WorkflowTests(unittest.TestCase):
    def test_export_is_real_workbook(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            s=Store(d); key=s.add({'url':'https://example.com/a','title':'React','company':'=1+1','description':'React'})
            s.update(key,'NEEDS_REVIEW',note='Awaiting review')
            wb=load_workbook(s.export()); ws=wb['Applications']
            self.assertEqual(ws['B2'].value,"'=1+1"); self.assertEqual(ws['E2'].value,'NEEDS_REVIEW')
            self.assertEqual(wb['History'].max_row,3)
            wb.close()
    def test_import_and_filter(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'jobs.json'; p.write_text(json.dumps([{'url':'https://example.com','company':'A','title':'React'}]))
            self.assertEqual(read_jobs(p)[0]['title'],'React')
        rows=[{'url':'https://example.com','title':'React Developer','location':'Remote US','description':'React','company':'A'}, {'url':'https://example.com/pl','title':'React Developer','location':'Poland Remote','description':'React','company':'A'}]
        with patch('jobpilot.sources.greenhouse',return_value=rows):
            found,errors=discover({'sources':[{'type':'greenhouse','board':'demo'}]})
        self.assertEqual(len(found),1); self.assertFalse(errors)
        self.assertEqual(found[0]['location'],'Poland Remote')
    def test_resume_writes_pdf_and_facts(self):
        with tempfile.TemporaryDirectory() as d:
            s=Store(d); key=s.add({'url':'https://example.com','company':'A','title':'React','description':'React'})
            pdf=prepare(s,key,PROFILE)
            self.assertTrue(pdf.read_bytes().startswith(b'%PDF'))
            self.assertEqual(read_json(pdf.parent/'facts.json')['experience'][0]['company'],'Demo Employer')
            self.assertEqual(s.get(key)['status'],'PREPARED')
    def test_real_browser_fills_and_receives_submission(self):
        try: from playwright.sync_api import sync_playwright
        except ImportError: self.skipTest('Install playwright for browser integration test')
        from jobpilot.browser import fill_common,fill_adapter,submit_once
        with tempfile.TemporaryDirectory() as d:
            server=make_server(0); thread=threading.Thread(target=server.serve_forever,daemon=True); thread.start()
            try:
                url=f'http://127.0.0.1:{server.server_port}'
                s=Store(d); key=s.add({'url':url,'company':'Local Demo','title':'React','description':'React'})
                pdf=prepare(s,key,PROFILE)
                adapter={'origin':url,'submit_selector':'button[type=submit]','success_selector':'#receipt','success_text':'Application received'}
                with sync_playwright() as p:
                    browser=p.chromium.launch(headless=True)
                    page=browser.new_page(); page.goto(url)
                    report=fill_common(page,PROFILE,pdf)
                    self.assertEqual(page.locator('#first').input_value(),'Demo')
                    self.assertEqual(page.locator('#email').input_value(),'demo@example.com')
                    self.assertEqual(page.locator('#country').input_value(),'PL')
                    self.assertFalse(page.locator('#consent').is_checked())
                    self.assertEqual(page.locator('#sponsor').input_value(),'')
                    self.assertTrue(any('(file)' in f for f in report['filled']))
                    page.locator('#consent').check(); page.locator('#sponsor').select_option('no')
                    s.claim(key)
                    self.assertTrue(submit_once(page,adapter,s,key))
                    self.assertEqual(s.get(key)['status'],'SUBMITTED'); self.assertFalse(s.can_apply(key))
                    browser.close()
            finally: server.shutdown(); server.server_close(); thread.join()

if __name__=='__main__': unittest.main()
